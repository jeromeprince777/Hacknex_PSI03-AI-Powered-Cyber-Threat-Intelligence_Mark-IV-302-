"""
Generate ThreatWeave HackNex 2026 Presentation Deck (.pptx)
Follows professional 16:9 widescreen layout, dark SOC theme, 
generous whitespace, embedded annotated screenshots, and speaker notes.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Paths
BASE_DIR = r"c:\Users\ffgam\OneDrive\Desktop\Hacknex"
ASSETS_DIR = os.path.join(BASE_DIR, "assets", "annotated")
OUTPUT_PPTX = os.path.join(BASE_DIR, "ThreatWeave_HackNex_Presentation.pptx")

# Color Palette (Dark Theme / Cybersecurity SOC)
COLOR_BG = RGBColor(11, 15, 25)          # Deep Navy/Midnight (#0B0F19)
COLOR_CARD = RGBColor(17, 24, 39)        # Dark Slate Card (#111827)
COLOR_CARD_BORDER = RGBColor(31, 41, 55) # Border (#1F2937)
COLOR_TEXT_PRIMARY = RGBColor(248, 250, 252) # White (#F8FAFC)
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)   # Slate Gray (#94A3B8)
COLOR_CYAN = RGBColor(6, 182, 212)       # Cyan (#06B6D4)
COLOR_EMERALD = RGBColor(16, 185, 129)   # Emerald (#10B981)
COLOR_ROSE = RGBColor(244, 63, 94)       # Rose (#F43F5E)
COLOR_AMBER = RGBColor(245, 158, 11)     # Amber (#F59E0B)
COLOR_PURPLE = RGBColor(168, 85, 247)    # Purple (#A855F7)
COLOR_BLUE = RGBColor(59, 130, 246)      # Blue (#3B82F6)

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def add_header(slide, title_text, category_text="THREATWEAVE // HNX26PSI03"):
    # Category Pill / Super-header
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.name = "Segoe UI"
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_CYAN

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_PRIMARY

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    return shape

def add_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    # Large Title Card
    c1 = add_card(s1, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1), bg_color=RGBColor(15, 23, 42), border_color=COLOR_CYAN)
    
    tb1 = s1.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.7), Inches(3.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "HACKNEX 2026 // HNX26PSI03"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_after = Pt(14)
    
    p2 = tf1.add_paragraph()
    p2.text = "THREATWEAVE"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_PRIMARY
    p2.space_after = Pt(8)
    
    p3 = tf1.add_paragraph()
    p3.text = "AI-Powered Cyber Threat Intelligence & Multi-Stage Attack Chain Reconstruction"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(20)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.space_after = Pt(24)
    
    p4 = tf1.add_paragraph()
    p4.text = "• Evidence-First Correlation Engine  • Deterministic 4-Factor Risk Index (89/100)\n• Real-Time Promotion Gate (0 FP on Benign)  • Bounded Anti-Hallucination Copilot"
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(14)
    p4.font.color.rgb = COLOR_EMERALD

    add_speaker_notes(s1, 
        "Good morning / afternoon, respected judges. Welcome to the presentation of ThreatWeave, "
        "developed specifically for HackNex 2026 Problem Statement HNX26PSI03: AI-Powered Cyber Threat Intelligence. "
        "ThreatWeave is an evidence-first security platform that solves the hardest challenge in modern SOC operations: "
        "connecting fragmented security logs across users, endpoints, IPs, and applications into chronological attack chains, "
        "while guaranteeing 100% telemetry grounding with zero generative hallucinations."
    )

    # =========================================================================
    # SLIDE 2: Problem Statement & Key Rules
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The PS03 Mandate & Core Architectural Rules")

    # Left Column: The Challenge
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb2_left = s2.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.6))
    tf2_l = tb2_left.text_frame
    tf2_l.word_wrap = True
    
    p = tf2_l.paragraphs[0]
    p.text = "THE PROBLEM: DISCONNECTED LOGS"
    p.font.name = "Segoe UI"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_ROSE
    p.space_after = Pt(10)
    
    bullets_l = [
        "Millions of logs collect daily from Windows, VPNs, EDR, and firewalls.",
        "Each log by itself appears normal or routine (login, file read, USB insert).",
        "Point-detection SIEMs flood analysts with isolated, uncorrelated alerts.",
        "Traditional AI tools hallucinate fake timestamps and ungrounded threat indicators."
    ]
    for b in bullets_l:
        p_b = tf2_l.add_paragraph()
        p_b.text = f"•  {b}"
        p_b.font.name = "Segoe UI"; p_b.font.size = Pt(13); p_b.font.color.rgb = COLOR_TEXT_MUTED
        p_b.space_after = Pt(8)

    # Right Column: The 2 Golden Rules
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb2_right = s2.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.6))
    tf2_r = tb2_right.text_frame
    tf2_r.word_wrap = True
    
    p = tf2_r.paragraphs[0]
    p.text = "AUTHORITATIVE PS03 EVALUATION RULES"
    p.font.name = "Segoe UI"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_AMBER
    p.space_after = Pt(10)

    rules_r = [
        ("RULE 1: FALSE POSITIVE LIMIT", "More alerts ≠ better score. Raising too many false alarms on clean logs caps your score. System must stay quiet on benign data.", COLOR_EMERALD),
        ("RULE 2: PROOF FOR EVERY STAGE", "Every stage of the attack must have proof. An unsupported claim doesn't count. Every promoted stage must cite real telemetry event IDs.", COLOR_CYAN),
        ("THREATWEAVE ARCHITECTURE", "Deterministic correlation engine + 2-stage promotion gate + bounded evidence pack + strict schema verifier.", COLOR_PURPLE)
    ]
    for title, desc, col in rules_r:
        p_t = tf2_r.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(13); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf2_r.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(12); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(8)

    add_speaker_notes(s2, 
        "In HackNex PS03, judges specifically emphasize two non-negotiable rules. "
        "Rule 1: More alerts do not equal a better score; false alarms cap your evaluation. "
        "Rule 2: Every single attack stage must have concrete proof; an unsupported claim receives zero points. "
        "ThreatWeave was engineered from day one to enforce these constraints mathematically."
    )

    # =========================================================================
    # SLIDE 3: System Pipeline & Data Architecture
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "End-to-End Pipeline & Data Flow Architecture")

    steps = [
        ("1. INGESTION", "Normalizes Windows, VPN, and EDR logs into CanonicalEvent-v1 schema with quarantine safety.", COLOR_BLUE),
        ("2. ATOMIC SIGNALS", "Declarative YAML rules & entity novelty baselines detect atomic anomalies (auth novelty, sensitive files).", COLOR_CYAN),
        ("3. CORRELATION", "TemporalSequenceCorrelator groups signals by primary actor across sliding 3600-second windows.", COLOR_PURPLE),
        ("4. PROMOTION GATE", "Enforces ≥3 events, ≥2 stages, and monotonic chronology. Suppressed noise logged to audit ledger.", COLOR_EMERALD),
        ("5. RISK ENGINE", "Computes explainable 4-factor score: 35% Signal + 30% Sequence + 20% Novelty + 15% Anomaly.", COLOR_AMBER),
        ("6. BOUNDED COPILOT", "Synthesizes analyst briefing strictly from bounded evidence pack. Rejects ungrounded tokens.", COLOR_ROSE),
    ]

    card_w = Inches(3.7)
    card_h = Inches(2.4)
    positions = [
        (Inches(0.8), Inches(1.6)), (Inches(4.8), Inches(1.6)), (Inches(8.8), Inches(1.6)),
        (Inches(0.8), Inches(4.3)), (Inches(4.8), Inches(4.3)), (Inches(8.8), Inches(4.3))
    ]

    for i, (title, desc, col) in enumerate(steps):
        x, y = positions[i]
        add_card(s3, x, y, card_w, card_h)
        tb = s3.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = col
        p.space_after = Pt(8)
        
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_speaker_notes(s3,
        "Here is the data pipeline running under the hood. Raw JSON lines enter through our normalizer. "
        "Atomic signals are extracted and clustered by actor. Then, our promotion gate checks if at least 3 events "
        "and 2 distinct attack stages exist. If so, an Incident is promoted and scored. If not, it is quietly recorded "
        "in our immutable suppression ledger. The AI Copilot only ever sees the promoted incident's bounded evidence pack."
    )

    # =========================================================================
    # SLIDE 4: Walkthrough 1 — Executive Console & Risk Meter
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Element Walkthrough: Executive Console & 4-Factor Risk Meter")

    # Left: Explanation Cards
    add_card(s4, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
    tb4 = s4.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
    tf4 = tb4.text_frame; tf4.word_wrap = True
    
    p = tf4.paragraphs[0]
    p.text = "CORE UI ELEMENTS HIGHLIGHTED"
    p.font.name = "Segoe UI"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_CYAN
    p.space_after = Pt(10)

    items4 = [
        ("[1] Actor Header", "Attributed to user-42 on device-17 across 4 attack stages and 7 correlated events.", COLOR_CYAN),
        ("[2] Radial Risk Meter", "Calculates composite Risk Index: 89/100 (HIGH) with 94% statistical confidence.", COLOR_ROSE),
        ("[3] 4-Factor Formula", "100 × (0.35 × Signal + 0.30 × Sequence + 0.20 × Novelty + 0.15 × Anomaly). Transparent and exposed.", COLOR_AMBER),
        ("[4] Action Playbook", "Approved dynamic response catalog: isolate-device, disable-session, review-resource-access.", COLOR_EMERALD)
    ]
    for tag, desc, col in items4:
        p_t = tf4.add_paragraph()
        p_t.text = tag
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(13); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf4.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(6)

    # Right: Annotated Screenshot
    img_path_1 = os.path.join(ASSETS_DIR, "annotated_01_executive_console.png")
    if os.path.exists(img_path_1):
        s4.shapes.add_picture(img_path_1, Inches(5.8), Inches(1.5), width=Inches(6.8))

    add_speaker_notes(s4,
        "Looking at the top of the Investigation Console: Callout 1 shows primary actor attribution to user-42. "
        "Callout 2 displays our circular radial risk meter showing an 89/100 score. "
        "Callout 3 exposes the exact 4-factor equation with individual weights, eliminating AI black-box scoring. "
        "Callout 4 highlights our approved action playbook with containment buttons."
    )

    # =========================================================================
    # SLIDE 5: Walkthrough 2 — Chronological Attack Timeline
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Element Walkthrough: Chronological Attack Timeline")

    add_card(s5, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
    tb5 = s5.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
    tf5 = tb5.text_frame; tf5.word_wrap = True
    
    p = tf5.paragraphs[0]
    p.text = "ATTACK RECONSTRUCTION STAGES"
    p.font.name = "Segoe UI"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_AMBER
    p.space_after = Pt(10)

    items5 = [
        ("[1] Stage 1: Initial Access", "09:41:12 UTC — VPN login from Bucharest, Romania (203.0.113.10) triggering AUTH_NEW_LOCATION (evt-003).", COLOR_CYAN),
        ("[2] Stage 2: Collection", "09:48:33 to 09:50:14 UTC — Access to payroll.xlsx and Project_Titan.pdf on corporate shares (evt-008, evt-009).", COLOR_PURPLE),
        ("[3] Stage 3: USB Transfer", "09:52:05 to 09:55:02 UTC — Removable storage mounted on device-17; files copied to Volume E:\\ (evt-011, 012, 013).", COLOR_AMBER),
        ("[4] Stage 4: Exfiltration", "09:58:45 UTC — 14.5 MB egress transfer to foreign destination 198.51.100.88 (evt-015).", COLOR_ROSE),
        ("Chronological Guarantee", "Strict timestamp ordering verified across all 15 events. 0% timestamp drift.", COLOR_EMERALD)
    ]
    for tag, desc, col in items5:
        p_t = tf5.add_paragraph()
        p_t.text = tag
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf5.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    img_path_2 = os.path.join(ASSETS_DIR, "annotated_02_attack_timeline.png")
    if os.path.exists(img_path_2):
        s5.shapes.add_picture(img_path_2, Inches(5.8), Inches(1.5), width=Inches(6.8))

    add_speaker_notes(s5,
        "Here is the core deliverable requested by PS03: Reconstructing what an attacker did, step by step. "
        "The conceptual example in the problem statement mentions: unusual login -> unusual file access -> USB copy. "
        "ThreatWeave reconstructs this exact chain across 4 discrete stages, strictly preserving chronological timestamps "
        "from 09:41 to 09:58 UTC. Analysts can click any stage card to filter the graph and open the evidence drawer."
    )

    # =========================================================================
    # SLIDE 6: Walkthrough 3 — Interactive SVG Entity Graph
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Element Walkthrough: Interactive SVG Entity-Relational Graph")

    add_card(s6, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
    tb6 = s6.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
    tf6 = tb6.text_frame; tf6.word_wrap = True
    
    p = tf6.paragraphs[0]
    p.text = "MULTI-RELATIONAL TOPOLOGY"
    p.font.name = "Segoe UI"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_CYAN
    p.space_after = Pt(10)

    items6 = [
        ("[1] Actor & Device Core Hub", "user-42 connected to device-17 via OCCURRED_ON relationships.", COLOR_CYAN),
        ("[2] Ingress IP Node", "203.0.113.10 mapped via FROM_IP link during VPN authentication.", COLOR_PURPLE),
        ("[3] Sensitive Resource Nodes", "payroll.xlsx and Project_Titan.pdf linked via ACCESSED_RESOURCE.", COLOR_EMERALD),
        ("[4] Removable Media Volume", "Volume:E:\\ connected via MOUNTED_USB to device-17.", COLOR_AMBER),
        ("[5] Egress Destination IP", "198.51.100.88 connected via CONNECTED_TO_IP with exfiltrated byte payload.", COLOR_ROSE),
        ("Dynamic Subgraph Filtering", "Built with NetworkX; supports zoom, drag, and stage-specific focus.", COLOR_BLUE)
    ]
    for tag, desc, col in items6:
        p_t = tf6.add_paragraph()
        p_t.text = tag
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf6.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    img_path_3 = os.path.join(ASSETS_DIR, "annotated_03_entity_graph.png")
    if os.path.exists(img_path_3):
        s6.shapes.add_picture(img_path_3, Inches(5.8), Inches(1.5), width=Inches(6.8))

    add_speaker_notes(s6,
        "This is our interactive SVG Entity Graph, powered by NetworkX on the backend. "
        "It connects 28 nodes: users, workstations, ingress IPs, sensitive file shares, USB volumes, and egress IPs. "
        "Judges can visibly see how user-42 connects to device-17, mounts Volume E:, accesses the confidential files, "
        "and pushes data out to 198.51.100.88. Clicking any stage dims unrelated nodes."
    )

    # =========================================================================
    # SLIDE 7: Walkthrough 4 — Forensic Evidence Drawer
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Element Walkthrough: Forensic Evidence Drawer (Rule #2 Proof)")

    add_card(s7, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
    tb7 = s7.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
    tf7 = tb7.text_frame; tf7.word_wrap = True
    
    p = tf7.paragraphs[0]
    p.text = "100% TELEMETRY GROUNDING"
    p.font.name = "Segoe UI"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_EMERALD
    p.space_after = Pt(10)

    items7 = [
        ("[1] Stage Forensic Anchor", "Exposes triggering detector (AUTH_NEW_LOCATION), base weight (0.55), and exact start/end timestamps.", COLOR_CYAN),
        ("[2] Concrete Event Card", "Cites verified event ID evt-003 with source vpn, event_type login_success, and action authenticate.", COLOR_EMERALD),
        ("[3] Raw JSON Payload", "Displays access mask, file path, source IP (203.0.113.10), and location (Bucharest, Romania).", COLOR_AMBER),
        ("Rule #2 Compliance", "'Every stage must have proof.' ThreatWeave enforces that every promoted stage must reference real event IDs.", COLOR_ROSE),
        ("Zero Hallucination", "Evaluators can compare the drawer directly to data/scenarios/attack_insider_usb.jsonl.", COLOR_BLUE)
    ]
    for tag, desc, col in items7:
        p_t = tf7.add_paragraph()
        p_t.text = tag
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf7.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    img_path_4 = os.path.join(ASSETS_DIR, "annotated_04_evidence_drawer.png")
    if os.path.exists(img_path_4):
        s7.shapes.add_picture(img_path_4, Inches(5.8), Inches(1.5), width=Inches(6.8))

    add_speaker_notes(s7,
        "PS03 Key Rule #2 states: 'Every stage of the attack must have proof. An unsupported claim doesn't count.' "
        "Here is our Forensic Evidence Drawer. When an evaluator clicks Stage 1, this drawer slides open. "
        "It cites the exact event ID evt-003, displaying the raw VPN payload, the Bucharest IP, and the exact timestamp. "
        "There are zero synthetic or inferred stages in ThreatWeave."
    )

    # =========================================================================
    # SLIDE 8: Walkthrough 5 — Bounded Copilot & Verifier Guard
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Element Walkthrough: Bounded Copilot & Anti-Hallucination Guard")

    add_card(s8, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
    tb8 = s8.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
    tf8 = tb8.text_frame; tf8.word_wrap = True
    
    p = tf8.paragraphs[0]
    p.text = "GROUNDED INCIDENT EXPLANATION"
    p.font.name = "Segoe UI"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_PURPLE
    p.space_after = Pt(10)

    items8 = [
        ("[1] Bounded Narrative", "Explains why this is an attack: correlates 7 events across 4 stages with factor weights (0.71 Sig, 1.00 Seq).", COLOR_CYAN),
        ("[2] Verification Passed Seal", "Passes CopilotOutputVerifier schema and set-membership checks against stored incident.", COLOR_EMERALD),
        ("[3] Grounded Event Links", "Clickable badges for evt-003, evt-008, evt-012 linking directly to forensic telemetry.", COLOR_BLUE),
        ("[4] Negative Test Sandbox", "Built-in evaluator buttons to inject fake event ID (evt-fake-9999) or unauthorized action.", COLOR_ROSE),
        ("Real-Time Rejection", "Injecting hallucinated tokens instantly triggers verifier rejection before display.", COLOR_AMBER)
    ]
    for tag, desc, col in items8:
        p_t = tf8.add_paragraph()
        p_t.text = tag
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf8.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    img_path_5 = os.path.join(ASSETS_DIR, "annotated_05_copilot_verified.png")
    if os.path.exists(img_path_5):
        s8.shapes.add_picture(img_path_5, Inches(5.8), Inches(1.5), width=Inches(6.8))

    add_speaker_notes(s8,
        "Many teams connect a generic LLM that fabricates indicators. In ThreatWeave, our Copilot operates strictly "
        "on a Bounded Evidence Pack. Callout 2 shows the green verification seal. "
        "Even better: Callout 4 highlights our built-in stress-test sandbox. During evaluation, judges can click "
        "'Inject Hallucinated Event evt-fake-9999', and watch our verifier immediately intercept and reject the output in real time!"
    )

    # =========================================================================
    # SLIDE 9: Walkthrough 6 — Replay Lab & Multi-Scenario Benchmark
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Element Walkthrough: Replay Lab & Multi-Scenario Benchmarking")

    add_card(s9, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
    tb9 = s9.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
    tf9 = tb9.text_frame; tf9.word_wrap = True
    
    p = tf9.paragraphs[0]
    p.text = "OFFICIAL BENCHMARK SCENARIOS"
    p.font.name = "Segoe UI"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_EMERALD
    p.space_after = Pt(10)

    items9 = [
        ("[1] Canonical Attack", "attack_insider_usb.jsonl (15 events) -> 1 Incident Promoted, 4 Stages, 89 Risk.", COLOR_CYAN),
        ("[2] Benign Admin Noise", "benign_admin.jsonl (10 events) -> 0 Incidents Promoted, 2 Suppressed, 100% FP quiet.", COLOR_EMERALD),
        ("[3] Partial Attack Chain", "attack_partial.jsonl (7 events) -> 1 Incident Promoted, 2 Stages, Risk gracefully degrades to 76.", COLOR_AMBER),
        ("[4] Benchmark Ledger", "Permanent table comparing records, signals, incidents, and promotion status.", COLOR_PURPLE),
        ("Single-Click Replay", "Allows evaluators to re-run any dataset from a clean slate in under 2 seconds.", COLOR_BLUE)
    ]
    for tag, desc, col in items9:
        p_t = tf9.add_paragraph()
        p_t.text = tag
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf9.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    img_path_6 = os.path.join(ASSETS_DIR, "annotated_06_replay_lab.png")
    if os.path.exists(img_path_6):
        s9.shapes.add_picture(img_path_6, Inches(5.8), Inches(1.5), width=Inches(6.8))

    add_speaker_notes(s9,
        "In the Replay Lab, evaluators can test reproducibility. We preloaded 3 distinct scenarios: "
        "The Canonical Attack, the Benign Administrator negative control, and a Partial Attack. "
        "Clicking 'Replay Dataset' re-runs the entire pipeline end-to-end. "
        "Notice that on benign_admin.jsonl, exactly 0 incidents are promoted, proving adherence to Key Rule #1."
    )

    # =========================================================================
    # SLIDE 10: Walkthrough 7 — False-Positive Suppression Audit Ledger
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Element Walkthrough: False-Positive Suppression Audit Ledger")

    add_card(s10, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.4))
    tb10 = s10.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(5.0))
    tf10 = tb10.text_frame; tf10.word_wrap = True
    
    p = tf10.paragraphs[0]
    p.text = "SILENCE WITHOUT BLIND SPOTS"
    p.font.name = "Segoe UI"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_PURPLE
    p.space_after = Pt(10)

    items10 = [
        ("[1] Suppression KPI Card", "Logs 2 suppressed clusters from benign dataset; prevents false alarms without dropping data.", COLOR_EMERALD),
        ("[2] Audited Cluster Records", "Records actor keys (user-admin-01, user-admin-02) and signals (UNUSUAL_FILE_ACCESS_TIME).", COLOR_PURPLE),
        ("Promotion Gate Reason", "Cluster labeled INSUFFICIENT_CHAIN_EVIDENCE (<3 events, <2 attack stages).", COLOR_AMBER),
        ("Compliance & Auditability", "Suppression is tracked in SQLite for regulatory review; nothing is silently discarded.", COLOR_CYAN),
        ("Rule #1 Adherence", "Guarantees quiet SOC queue during benign IT operations while logging every decision.", COLOR_BLUE)
    ]
    for tag, desc, col in items10:
        p_t = tf10.add_paragraph()
        p_t.text = tag
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = col
        p_d = tf10.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    img_path_7 = os.path.join(ASSETS_DIR, "annotated_07_suppression_ledger.png")
    if os.path.exists(img_path_7):
        s10.shapes.add_picture(img_path_7, Inches(5.8), Inches(1.5), width=Inches(6.8))

    add_speaker_notes(s10,
        "When clean logs enter ThreatWeave, what happens to isolated signals like an administrator working after hours? "
        "Instead of spamming the SOC triage queue or dropping the event into a black hole, our system routes it "
        "to this immutable Suppression Audit Ledger. It documents why the cluster was suppressed—specifically, "
        "INSUFFICIENT_CHAIN_EVIDENCE—maintaining full compliance while keeping the triage queue clean."
    )

    # =========================================================================
    # SLIDE 11: Dataset Provenance & Technology Audit
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Dataset Provenance & Complete Technology Disclosure")

    add_card(s11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb11_l = s11.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.6))
    tf11_l = tb11_l.text_frame; tf11_l.word_wrap = True

    p = tf11_l.paragraphs[0]
    p.text = "DATASET PROVENANCE DISCLOSURE"
    p.font.name = "Segoe UI"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_CYAN
    p.space_after = Pt(10)

    data_items = [
        ("attack_insider_usb.jsonl (15 events)", "Team-generated deterministic synthetic telemetry covering canonical 4-stage insider threat."),
        ("benign_admin.jsonl (10 events)", "Team-generated negative control telemetry modeling routine sysadmin activity."),
        ("attack_partial.jsonl (7 events)", "Team-generated partial scenario evaluating graceful risk score degradation."),
        ("malformed_logs.jsonl (8 records)", "Quarantine test fixture verifying schema validation and bad data handling."),
        ("Public Datasets Used:", "NONE. Zero external public datasets (DARPA/CICIDS) used. 100% reproducible committed JSONL.")
    ]
    for title, desc in data_items:
        p_t = tf11_l.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf11_l.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    add_card(s11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb11_r = s11.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.6))
    tf11_r = tb11_r.text_frame; tf11_r.word_wrap = True

    p = tf11_r.paragraphs[0]
    p.text = "TECHNOLOGY & ENVIRONMENT AUDIT"
    p.font.name = "Segoe UI"; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_EMERALD
    p.space_after = Pt(10)

    tech_items = [
        ("Backend Core:", "Python 3.11+ / FastAPI, Pydantic v2, NetworkX 3.4, NumPy 2.2, Pandas 2.2, SQLite."),
        ("Frontend Console:", "Next.js 16 (React 19), Tailwind CSS v4, Lucide Icons, Interactive SVG Entity Graph."),
        ("AI / ML Architecture:", "DeterministicCopilotProvider + CopilotOutputVerifier; RobustAnomalyScorer (z-excess)."),
        ("Network Independence:", "100% Air-Gapped / Offline. Zero external HTTP calls. Zero external API keys needed."),
        ("Test Verification:", "42 of 42 Automated Pytest Tests Passing (100%). Playwright E2E 8/8 Steps Passing.")
    ]
    for title, desc in tech_items:
        p_t = tf11_r.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.name = "Segoe UI"; p_t.font.size = Pt(12); p_t.font.bold = True; p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf11_r.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"; p_d.font.size = Pt(11); p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_after = Pt(4)

    add_speaker_notes(s11,
        "Under the HackNex 'Declare Your Resources' rule, we are completely transparent about our technology and data. "
        "All 4 scenario datasets are committed, static JSONL files with verified SHA-256 hashes. "
        "No external public datasets like DARPA or CICIDS are used in runtime. "
        "The system runs 100% offline, requires zero API keys, and has 42 automated tests passing with 100% pass rate."
    )

    # =========================================================================
    # SLIDE 12: Judge Presentation Blueprint & Live Demo Flow
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Judge Presentation Blueprint: 3-Minute Live Demo Flow")

    demo_steps = [
        ("MINUTE 1: THE RECONSTRUCTION", [
            "1. Open Investigation Console on inc-001.",
            "2. Highlight Circular Risk Meter (89/100) and explain the 4-factor formula.",
            "3. Walk across the 4 Attack Timeline stages: Initial Access -> Collection -> USB Transfer -> Exfiltration.",
            "4. Prove chronological monotonicity from 09:41 to 09:58 UTC."
        ], COLOR_CYAN),
        ("MINUTE 2: EVIDENCE & GRAPH", [
            "1. Click Stage 1: Watch the SVG Entity Graph highlight user-42 and device-17.",
            "2. Open Forensic Evidence Drawer: Point out real event ID evt-003 with raw VPN payload.",
            "3. Show Analyst Copilot: Highlight the green 'EVIDENCE VALIDATION PASSED' seal.",
            "4. Show Grounded Citations: Click evt-003 and evt-008 badges."
        ], COLOR_PURPLE),
        ("MINUTE 3: RULES & CONTROLS", [
            "1. Switch to Replay Lab: Click 'Replay Dataset' on benign_admin.jsonl.",
            "2. Show 0 Incidents Promoted, proving Key Rule #1 (False-positive limit).",
            "3. Open Suppression Ledger: Show the 2 audited clusters with reason INSUFFICIENT_CHAIN_EVIDENCE.",
            "4. Click 'Inject Fake Event': Show the Copilot Verifier instantly intercepting hallucinated tokens."
        ], COLOR_EMERALD)
    ]

    card_w = Inches(3.7)
    card_h = Inches(5.2)
    x_positions = [Inches(0.8), Inches(4.8), Inches(8.8)]

    for i, (col_title, bullet_list, col) in enumerate(demo_steps):
        x = x_positions[i]
        add_card(s12, x, Inches(1.6), card_w, card_h)
        tb = s12.shapes.add_textbox(x + Inches(0.2), Inches(1.8), card_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame; tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = col_title
        p.font.name = "Segoe UI"; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = col
        p.space_after = Pt(12)

        for b in bullet_list:
            p_b = tf.add_paragraph()
            p_b.text = b
            p_b.font.name = "Segoe UI"; p_b.font.size = Pt(11); p_b.font.color.rgb = COLOR_TEXT_MUTED
            p_b.space_after = Pt(8)

    add_speaker_notes(s12,
        "Here is our 3-minute live evaluation blueprint. "
        "In Minute 1, we show the canonical attack and the 4-factor risk meter. "
        "In Minute 2, we demonstrate Rule #2 by opening the evidence drawer and graph for Stage 1. "
        "In Minute 3, we execute the benign replay to prove Rule #1 false-positive suppression, and click "
        "the fake injection button to prove real-time verifier interception. "
        "Thank you! We are now ready for your questions."
    )

    prs.save(OUTPUT_PPTX)
    print(f"Presentation saved successfully to: {OUTPUT_PPTX}")

if __name__ == "__main__":
    build_deck()
